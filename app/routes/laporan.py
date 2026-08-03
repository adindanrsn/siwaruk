from flask import Blueprint, render_template, request, redirect, url_for, flash, send_file, jsonify, current_app
from flask_login import login_required, current_user
from app.extensions import db
from app.models.sale import Sale
from app.models.sale_detail import SaleDetail
from app.models.expense import Expense
from app.utils import get_active_business
from sqlalchemy import func
from datetime import datetime, timedelta, timezone as _tz
from zoneinfo import ZoneInfo
import calendar
import io
import openpyxl
import os
from openpyxl.styles import Font, Alignment
# PDF Export imports
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER

laporan_bp = Blueprint('laporan', __name__, url_prefix='/laporan')
_WIB = ZoneInfo('Asia/Jakarta')

def _check_owner_and_active_biz():
    active_biz = get_active_business()
    if not active_biz:
        flash('Pilih atau buat usaha terlebih dahulu.', 'warning')
        return None, redirect(url_for('usaha.kelola'))
    if current_user.role != 'owner':
        flash('Anda tidak memiliki akses ke halaman laporan.', 'danger')
        return None, redirect(url_for('dashboard.index'))
    return active_biz, None

def get_date_range(periode, start_str, end_str):
    now = datetime.now(_WIB)
    start_wib, end_wib = None, None
    
    if periode == 'hari_ini':
        start_wib = now.replace(hour=0, minute=0, second=0, microsecond=0)
        end_wib = now.replace(hour=23, minute=59, second=59, microsecond=999999)
    elif periode == 'minggu_ini':
        start_wib = (now - timedelta(days=now.weekday())).replace(hour=0, minute=0, second=0, microsecond=0)
        end_wib = (start_wib + timedelta(days=6)).replace(hour=23, minute=59, second=59, microsecond=999999)
    elif periode == 'bulan_ini':
        start_wib = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        last_day = calendar.monthrange(now.year, now.month)[1]
        end_wib = now.replace(day=last_day, hour=23, minute=59, second=59, microsecond=999999)
    elif periode == 'tahun_ini':
        start_wib = now.replace(month=1, day=1, hour=0, minute=0, second=0, microsecond=0)
        end_wib = now.replace(month=12, day=31, hour=23, minute=59, second=59, microsecond=999999)
    elif periode == 'custom' and start_str and end_str:
        try:
            st = datetime.strptime(start_str, '%Y-%m-%d').replace(tzinfo=_WIB)
            en = datetime.strptime(end_str, '%Y-%m-%d').replace(tzinfo=_WIB)
            start_wib = st.replace(hour=0, minute=0, second=0, microsecond=0)
            end_wib = en.replace(hour=23, minute=59, second=59, microsecond=999999)
        except ValueError:
            pass

    # Fallback to bulan_ini if invalid custom or missing
    if not start_wib or not end_wib:
        start_wib = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        last_day = calendar.monthrange(now.year, now.month)[1]
        end_wib = now.replace(day=last_day, hour=23, minute=59, second=59, microsecond=999999)
        periode = 'bulan_ini'

    return start_wib, end_wib, periode

def to_utc_naive(dt_wib):
    return dt_wib.astimezone(_tz.utc).replace(tzinfo=None)

def from_utc_naive(dt_utc):
    if dt_utc is None: return None
    if dt_utc.tzinfo is None:
        dt_utc = dt_utc.replace(tzinfo=_tz.utc)
    return dt_utc.astimezone(_WIB)

def get_report_data(active_biz, start_wib, end_wib, periode):
    start_utc = to_utc_naive(start_wib)
    end_utc = to_utc_naive(end_wib)

    # 1. Total Pemasukan
    sales = Sale.query.filter(
        Sale.business_id == active_biz.id,
        Sale.transaction_date >= start_utc,
        Sale.transaction_date <= end_utc
    ).all()
    total_pemasukan = sum(s.total for s in sales)

    # 2. Total Pengeluaran
    expenses = Expense.query.filter(
        Expense.business_id == active_biz.id,
        Expense.expense_date >= start_utc,
        Expense.expense_date <= end_utc
    ).all()
    total_pengeluaran = sum(e.amount for e in expenses)

    # 3. Laba Bersih
    laba_bersih = total_pemasukan - total_pengeluaran

    # 4. Top Products (Limit 5)
    top_products = db.session.query(
        SaleDetail.product_id,
        func.sum(SaleDetail.quantity).label('qty'),
        func.sum(SaleDetail.subtotal).label('subtotal')
    ).join(Sale, Sale.id == SaleDetail.sale_id).filter(
        Sale.business_id == active_biz.id,
        Sale.transaction_date >= start_utc,
        Sale.transaction_date <= end_utc
    ).group_by(SaleDetail.product_id).order_by(db.text('qty DESC')).limit(5).all()

    from app.models.product import Product
    top_products_list = []
    for pid, qty, subtotal in top_products:
        prod = Product.query.get(pid)
        if prod:
            top_products_list.append({
                'name': prod.name,
                'qty': qty,
                'subtotal': subtotal
            })

    # 5. Chart Data
    chart_pemasukan = {}
    chart_pengeluaran = {}
    
    # helper for bucketing
    def _add_to_dict(d, dt, val, grouping):
        if grouping == 'hour':
            lbl = f"{dt.hour:02d}.00"
        elif grouping == 'day_name': # minggu_ini
            days = ['Sen', 'Sel', 'Rab', 'Kam', 'Jum', 'Sab', 'Min']
            lbl = days[dt.weekday()]
        elif grouping == 'date_only': # bulan_ini
            lbl = str(dt.day)
        elif grouping == 'month_name': # tahun_ini
            months = ['Jan', 'Feb', 'Mar', 'Apr', 'Mei', 'Jun', 'Jul', 'Agt', 'Sep', 'Okt', 'Nov', 'Des']
            lbl = months[dt.month - 1]
        elif grouping == 'custom_day':
            lbl = dt.strftime('%d %b %Y')
        else: # custom_month
            months = ['Jan', 'Feb', 'Mar', 'Apr', 'Mei', 'Jun', 'Jul', 'Agt', 'Sep', 'Okt', 'Nov', 'Des']
            lbl = f"{months[dt.month - 1]} {dt.year}"
        d[lbl] = d.get(lbl, 0) + float(val)

    # Determine grouping
    if periode == 'hari_ini':
        grouping = 'hour'
        labels = [f"{i:02d}.00" for i in range(24)]
    elif periode == 'minggu_ini':
        grouping = 'day_name'
        labels = ['Sen', 'Sel', 'Rab', 'Kam', 'Jum', 'Sab', 'Min']
    elif periode == 'bulan_ini':
        grouping = 'date_only'
        last_day = calendar.monthrange(start_wib.year, start_wib.month)[1]
        labels = [str(i) for i in range(1, last_day + 1)]
    elif periode == 'tahun_ini':
        grouping = 'month_name'
        labels = ['Jan', 'Feb', 'Mar', 'Apr', 'Mei', 'Jun', 'Jul', 'Agt', 'Sep', 'Okt', 'Nov', 'Des']
    else: # custom
        days_diff = (end_wib - start_wib).days
        if days_diff <= 1:
            grouping = 'hour'
            labels = [f"{i:02d}.00" for i in range(24)]
        elif days_diff <= 62:
            grouping = 'custom_day'
            labels = []
            curr = start_wib
            while curr <= end_wib:
                labels.append(curr.strftime('%d %b %Y'))
                curr += timedelta(days=1)
        else:
            grouping = 'custom_month'
            labels = []
            curr = start_wib.replace(day=1)
            months_indo = ['Jan', 'Feb', 'Mar', 'Apr', 'Mei', 'Jun', 'Jul', 'Agt', 'Sep', 'Okt', 'Nov', 'Des']
            while curr <= end_wib:
                labels.append(f"{months_indo[curr.month - 1]} {curr.year}")
                # add 1 month
                nxt_month = curr.month % 12 + 1
                nxt_year = curr.year + (curr.month // 12)
                curr = curr.replace(year=nxt_year, month=nxt_month, day=1)
            
    for s in sales:
        dt_wib = from_utc_naive(s.transaction_date)
        _add_to_dict(chart_pemasukan, dt_wib, s.total, grouping)
        
    for e in expenses:
        dt_wib = from_utc_naive(e.expense_date)
        _add_to_dict(chart_pengeluaran, dt_wib, e.amount, grouping)

    data_pemasukan = []
    data_pengeluaran_list = []
    data_laba_bersih = []
    
    for lbl in labels:
        p = chart_pemasukan.get(lbl, 0)
        e = chart_pengeluaran.get(lbl, 0)
        data_pemasukan.append(p)
        data_pengeluaran_list.append(e)
        data_laba_bersih.append(p - e)

    return {
        'total_pemasukan': total_pemasukan,
        'total_pengeluaran': total_pengeluaran,
        'laba_bersih': laba_bersih,
        'top_products': top_products_list,
        'chart_labels': labels,
        'chart_pemasukan': data_pemasukan,
        'chart_pengeluaran': data_pengeluaran_list,
        'chart_laba_bersih': data_laba_bersih,
        'sales': sales,
        'expenses': expenses
    }

def format_periode_id(start_date, end_date, periode):
    bulan_full = ["", "Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember"]
    bulan_short = ["", "Jan", "Feb", "Mar", "Apr", "Mei", "Jun", "Jul", "Agu", "Sep", "Okt", "Nov", "Des"]

    if periode == 'hari_ini':
        return f"{start_date.day} {bulan_full[start_date.month]} {start_date.year}"
    elif periode == 'bulan_ini':
        return f"{bulan_full[start_date.month]} {start_date.year}"
    elif periode == 'tahun_ini':
        return f"{start_date.year}"
    else:
        if start_date.year == end_date.year and start_date.month == end_date.month:
            if start_date.day == end_date.day:
                return f"{start_date.day} {bulan_full[start_date.month]} {start_date.year}"
            return f"{start_date.day}–{end_date.day} {bulan_full[start_date.month]} {start_date.year}"
        else:
            return f"{start_date.day} {bulan_short[start_date.month]} {start_date.year} – {end_date.day} {bulan_short[end_date.month]} {end_date.year}"


@laporan_bp.route('/', methods=['GET'])
@login_required
def index():
    active_biz, err = _check_owner_and_active_biz()
    if err: return err

    periode = request.args.get('periode', 'bulan_ini')
    start_str = request.args.get('start_date')
    end_str = request.args.get('end_date')

    start_wib, end_wib, periode = get_date_range(periode, start_str, end_str)
    
    formatted_periode = format_periode_id(start_wib, end_wib, periode)
    data = get_report_data(active_biz, start_wib, end_wib, periode)

    from app.models.laporan_terkirim import LaporanTerkirim
    current_report_terkirim = LaporanTerkirim.query.filter_by(
        business_id=active_biz.id,
        start_date=start_wib.date(),
        end_date=end_wib.date()
    ).order_by(LaporanTerkirim.submitted_at.desc()).first()

    riwayat_laporan_terkirim = LaporanTerkirim.query.filter_by(
        business_id=active_biz.id
    ).order_by(LaporanTerkirim.submitted_at.desc()).limit(10).all()

    if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
        # Remove full objects before returning JSON
        data.pop('sales', None)
        data.pop('expenses', None)
        return jsonify({
            'status': 'success',
            'periode': periode,
            'start_date': start_wib.strftime('%Y-%m-%d'),
            'end_date': end_wib.strftime('%Y-%m-%d'),
            'formatted_periode': formatted_periode,
            'data': data,
            'sent_status': current_report_terkirim.status if current_report_terkirim else None
        })

    return render_template(
        'laporan/index.html',
        active_business=active_biz,
        periode=periode,
        start_date=start_wib.strftime('%Y-%m-%d'),
        end_date=end_wib.strftime('%Y-%m-%d'),
        formatted_periode=formatted_periode,
        data=data,
        current_report_terkirim=current_report_terkirim,
        riwayat_laporan_terkirim=riwayat_laporan_terkirim
    )

def _generate_daily_rekap(start_wib, end_wib, data):
    rekap = {}
    curr = start_wib.replace(hour=0, minute=0, second=0, microsecond=0)
    end_date = end_wib.replace(hour=0, minute=0, second=0, microsecond=0)
    
    while curr <= end_date:
        key = curr.strftime('%Y-%m-%d')
        rekap[key] = {'pemasukan': 0.0, 'pengeluaran': 0.0, 'date_obj': curr}
        curr += timedelta(days=1)
        
    for s in data.get('sales', []):
        dt_wib = from_utc_naive(s.transaction_date)
        key = dt_wib.strftime('%Y-%m-%d')
        if key in rekap:
            rekap[key]['pemasukan'] += float(s.total)
            
    for e in data.get('expenses', []):
        dt_wib = from_utc_naive(e.expense_date)
        key = dt_wib.strftime('%Y-%m-%d')
        if key in rekap:
            rekap[key]['pengeluaran'] += float(e.amount)
            
    sorted_rekap = []
    for key in sorted(rekap.keys()):
        val = rekap[key]
        p = val['pemasukan']
        e = val['pengeluaran']
        laba = p - e
        sorted_rekap.append({
            'tanggal': val['date_obj'].strftime('%d %b %Y'),
            'pemasukan': p,
            'pengeluaran': e,
            'laba_bersih': laba
        })
    return sorted_rekap

def _create_header_footer_watermark(start_wib, end_wib, active_biz):
    def on_page(canvas, doc):
        canvas.saveState()
        
        # --- WATERMARK ---
        canvas.setFont('Helvetica-Bold', 60)
        canvas.setFillColorRGB(0.9, 0.9, 0.9, alpha=0.5)
        canvas.translate(297.5, 420.5)
        canvas.rotate(45)
        canvas.drawCentredString(0, 0, "SIWARUK")
        canvas.rotate(-45)
        canvas.translate(-297.5, -420.5)
        
        # --- FOOTER ---
        canvas.setFont('Helvetica', 9)
        canvas.setFillColorRGB(0.5, 0.5, 0.5)
        canvas.drawString(30, 20, "Dibuat menggunakan aplikasi Siwaruk")
        canvas.drawRightString(A4[0] - 30, 20, f"Halaman {doc.page}")
        
        canvas.restoreState()
    return on_page

def _add_excel_header(ws, start_wib, end_wib, active_biz, report_title, add_logo=False):
    months = ["", "Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember"]
    
    d1 = f"{start_wib.day} {months[start_wib.month]} {start_wib.year}"
    d2 = f"{end_wib.day} {months[end_wib.month]} {end_wib.year}"
    period_text = d1 if d1 == d2 else f"{d1} – {d2}"
    
    export_wib = datetime.now(_WIB)
    export_d = f"{export_wib.day} {months[export_wib.month]} {export_wib.year}"
    export_t = export_wib.strftime('%H.%M')

    if add_logo:
        ws.column_dimensions['A'].width = 16
        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 20
        
        ws.append(["", report_title])
        ws.append(["", active_biz.business_name])
        ws.append(["", ""])
        ws.append(["", f"Periode: {period_text}"])
        ws.append(["", f"Tanggal Export: {export_d} • {export_t} WIB"])
        ws.append([])
        
        for r in (1, 2, 4, 5):
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
            ws.cell(row=r, column=2).alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            
        ws['B1'].font = Font(bold=True, size=15)
        ws['B2'].font = Font(bold=True, size=13)
        ws['B4'].font = Font(size=11)
        ws['B5'].font = Font(size=9, italic=True)
        
        logo_path = os.path.join(current_app.root_path, 'static', 'img', 'logo-siwaruk.png')
        if os.path.exists(logo_path):
            try:
                from openpyxl.drawing.image import Image as ExcelImage
                img = ExcelImage(logo_path)
                aspect = img.width / img.height
                img.height = 55
                img.width = 55 * aspect
                img.anchor = 'A1'
                ws.add_image(img)
            except:
                pass
    else:
        ws.append([report_title])
        ws.append([active_biz.business_name])
        ws.append([])
        ws.append([f"Periode: {period_text}"])
        ws.append([f"Tanggal Export: {export_d} • {export_t} WIB"])
        ws.append([])
        
        for r in (1, 2, 4, 5):
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
            ws.cell(row=r, column=1).alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            
        ws['A1'].font = Font(bold=True, size=15)
        ws['A2'].font = Font(bold=True, size=13)
        ws['A4'].font = Font(size=11)
        ws['A5'].font = Font(size=9, italic=True)

@laporan_bp.route('/export/pdf', methods=['GET'])
@login_required
def export_pdf():
    active_biz, err = _check_owner_and_active_biz()
    if err: return err

    periode = request.args.get('periode', 'bulan_ini')
    start_wib, end_wib, _ = get_date_range(periode, request.args.get('start_date'), request.args.get('end_date'))
    data = get_report_data(active_biz, start_wib, end_wib, periode)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=40, bottomMargin=50)
    elements = []
    styles = getSampleStyleSheet()
    
    # --- PDF HEADER IN ELEMENTS ---
    logo_path = os.path.join(current_app.root_path, 'static', 'img', 'logo-siwaruk.png')
    if os.path.exists(logo_path):
        try:
            from reportlab.platypus import Image as RLImage
            img = RLImage(logo_path, width=60, height=60, kind='proportional')
            img.hAlign = 'CENTER'
            elements.append(img)
            elements.append(Spacer(1, 10))
        except:
            pass
            
    title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], alignment=TA_CENTER, fontSize=16, spaceAfter=8)
    elements.append(Paragraph(f"Laporan Keuangan {active_biz.business_name}", title_style))
    
    months = ["", "Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember"]
    d1 = f"{start_wib.day} {months[start_wib.month]} {start_wib.year}"
    d2 = f"{end_wib.day} {months[end_wib.month]} {end_wib.year}"
    period_text = d1 if d1 == d2 else f"{d1} – {d2}"
    
    period_style = ParagraphStyle('PeriodStyle', parent=styles['Normal'], alignment=TA_CENTER, fontSize=12, spaceAfter=4, leading=14)
    elements.append(Paragraph(f"Periode<br/>{period_text}", period_style))
    
    export_wib = datetime.now(_WIB)
    export_d = f"{export_wib.day} {months[export_wib.month]} {export_wib.year}"
    export_t = export_wib.strftime('%H.%M')
    export_style = ParagraphStyle('ExportStyle', parent=styles['Normal'], alignment=TA_CENTER, fontSize=10, textColor=colors.dimgrey, leading=12)
    elements.append(Paragraph(f"Tanggal Export<br/>{export_d} &bull; {export_t} WIB", export_style))
    
    elements.append(Spacer(1, 20))

    # Ringkasan
    summary_data = [
        ['Total Pemasukan', f"Rp {int(data['total_pemasukan']):,}".replace(',', '.')],
        ['Total Pengeluaran', f"Rp {int(data['total_pengeluaran']):,}".replace(',', '.')],
        ['Laba Bersih', f"Rp {int(data['laba_bersih']):,}".replace(',', '.')],
    ]
    t = Table(summary_data, colWidths=[200, 200])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.white),
        ('GRID', (0,0), (-1,-1), 1, colors.black),
        ('FONTNAME', (0,0), (-1,-1), 'Helvetica-Bold'),
        ('PADDING', (0,0), (-1,-1), 6)
    ]))
    elements.append(t)
    
    # Rekap Harian
    daily_rekap = _generate_daily_rekap(start_wib, end_wib, data)
    if daily_rekap:
        elements.append(Spacer(1, 20))
        elements.append(Paragraph("Rekap Harian", styles['Heading3']))
        
        rekap_table_data = [['Tanggal', 'Total Pemasukan', 'Total Pengeluaran', 'Laba Bersih']]
        for row in daily_rekap:
            rekap_table_data.append([
                row['tanggal'],
                f"Rp {int(row['pemasukan']):,}".replace(',', '.'),
                f"Rp {int(row['pengeluaran']):,}".replace(',', '.'),
                f"Rp {int(row['laba_bersih']):,}".replace(',', '.')
            ])
            
        rt = Table(rekap_table_data, colWidths=[120, 130, 130, 130])
        rt.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.lightgrey),
            ('GRID', (0,0), (-1,-1), 1, colors.black),
            ('PADDING', (0,0), (-1,-1), 6)
        ]))
        elements.append(rt)

    # Top Products
    if data['top_products']:
        elements.append(Spacer(1, 20))
        elements.append(Paragraph("5 Produk Terlaris", styles['Heading3']))
        prod_data = [['Nama Produk', 'Terjual', 'Total Penjualan']]
        for p in data['top_products']:
            prod_data.append([p['name'], str(p['qty']), f"Rp {int(p['subtotal']):,}".replace(',', '.')])
        
        pt = Table(prod_data, colWidths=[250, 80, 150])
        pt.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.lightgrey),
            ('GRID', (0,0), (-1,-1), 1, colors.black),
            ('PADDING', (0,0), (-1,-1), 6)
        ]))
        elements.append(pt)

    on_page_func = _create_header_footer_watermark(start_wib, end_wib, active_biz)
    doc.build(elements, onFirstPage=on_page_func, onLaterPages=on_page_func)
    buffer.seek(0)
    
    filename = f"Laporan_{start_wib.strftime('%Y%m%d')}_{end_wib.strftime('%Y%m%d')}.pdf"
    return send_file(buffer, as_attachment=True, download_name=filename, mimetype='application/pdf')

@laporan_bp.route('/export/excel', methods=['GET'])
@login_required
def export_excel():
    active_biz, err = _check_owner_and_active_biz()
    if err: return err

    periode = request.args.get('periode', 'bulan_ini')
    start_wib, end_wib, _ = get_date_range(periode, request.args.get('start_date'), request.args.get('end_date'))
    data = get_report_data(active_biz, start_wib, end_wib, periode)

    wb = openpyxl.Workbook()
    
    # Sheet 1: Ringkasan
    ws_sum = wb.active
    ws_sum.title = "Ringkasan"
    _add_excel_header(ws_sum, start_wib, end_wib, active_biz, "Laporan Keuangan", add_logo=True)
    ws_sum.append(["Keterangan", "Total (Rp)"])
    ws_sum.append(["Total Pemasukan", float(data['total_pemasukan'])])
    ws_sum.append(["Total Pengeluaran", float(data['total_pengeluaran'])])
    ws_sum.append(["Laba Bersih", float(data['laba_bersih'])])

    # Sheet 2: Rekap Harian
    daily_rekap = _generate_daily_rekap(start_wib, end_wib, data)
    ws_rekap = wb.create_sheet("Rekap Harian")
    _add_excel_header(ws_rekap, start_wib, end_wib, active_biz, "Rekap Harian")
    ws_rekap.append(["Tanggal", "Total Pemasukan (Rp)", "Total Pengeluaran (Rp)", "Laba Bersih (Rp)"])
    for row in daily_rekap:
        ws_rekap.append([row['tanggal'], float(row['pemasukan']), float(row['pengeluaran']), float(row['laba_bersih'])])

    # Sheet 3: Produk Terlaris
    ws_prod = wb.create_sheet("Produk Terlaris")
    _add_excel_header(ws_prod, start_wib, end_wib, active_biz, "Produk Terlaris")
    ws_prod.append(["Nama Produk", "Terjual", "Total Penjualan (Rp)"])
    for p in data['top_products']:
        ws_prod.append([p['name'], p['qty'], float(p['subtotal'])])

    # Sheet 4: Riwayat Pemasukan
    ws_in = wb.create_sheet("Riwayat Pemasukan")
    _add_excel_header(ws_in, start_wib, end_wib, active_biz, "Riwayat Pemasukan")
    ws_in.append(["No Invoice", "Tanggal", "Metode Pembayaran", "Total (Rp)"])
    for s in data['sales']:
        dt = from_utc_naive(s.transaction_date).strftime('%Y-%m-%d %H:%M')
        ws_in.append([s.invoice_number, dt, s.payment_method, float(s.total)])

    # Sheet 5: Riwayat Pengeluaran
    ws_ex = wb.create_sheet("Riwayat Pengeluaran")
    _add_excel_header(ws_ex, start_wib, end_wib, active_biz, "Riwayat Pengeluaran")
    ws_ex.append(["Keperluan", "Kategori", "Tanggal", "Total (Rp)"])
    for e in data['expenses']:
        dt = from_utc_naive(e.expense_date).strftime('%Y-%m-%d %H:%M')
        cat = e.category_name or (e.expense_category.name if e.expense_category else "-")
        ws_ex.append([e.description, cat, dt, float(e.amount)])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    filename = f"Laporan_{start_wib.strftime('%Y%m%d')}_{end_wib.strftime('%Y%m%d')}.xlsx"
    return send_file(buffer, as_attachment=True, download_name=filename, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ─────────────────────────────────────────────────────────────
#  HELPER: Buat elemen PDF (reusable untuk export dan kirim)
# ─────────────────────────────────────────────────────────────
def _build_pdf_elements(active_biz, start_wib, end_wib, periode, data):
    """Kembalikan (elements, on_page_func) untuk membangun PDF ReportLab."""
    elements = []
    styles = getSampleStyleSheet()

    logo_path = os.path.join(current_app.root_path, 'static', 'img', 'logo-siwaruk.png')
    if os.path.exists(logo_path):
        try:
            from reportlab.platypus import Image as RLImage
            img = RLImage(logo_path, width=60, height=60, kind='proportional')
            img.hAlign = 'CENTER'
            elements.append(img)
            elements.append(Spacer(1, 10))
        except Exception:
            pass

    title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'],
                                  alignment=TA_CENTER, fontSize=16, spaceAfter=8)
    elements.append(Paragraph(f"Laporan Keuangan {active_biz.business_name}", title_style))

    months = ["", "Januari", "Februari", "Maret", "April", "Mei", "Juni",
              "Juli", "Agustus", "September", "Oktober", "November", "Desember"]
    d1 = f"{start_wib.day} {months[start_wib.month]} {start_wib.year}"
    d2 = f"{end_wib.day} {months[end_wib.month]} {end_wib.year}"
    period_text = d1 if d1 == d2 else f"{d1} \u2013 {d2}"

    period_style = ParagraphStyle('PeriodStyle', parent=styles['Normal'],
                                   alignment=TA_CENTER, fontSize=12, spaceAfter=4, leading=14)
    elements.append(Paragraph(f"Periode<br/>{period_text}", period_style))

    export_wib = datetime.now(_WIB)
    export_d = f"{export_wib.day} {months[export_wib.month]} {export_wib.year}"
    export_t = export_wib.strftime('%H.%M')
    export_style = ParagraphStyle('ExportStyle', parent=styles['Normal'],
                                   alignment=TA_CENTER, fontSize=10,
                                   textColor=colors.dimgrey, leading=12)
    elements.append(Paragraph(f"Tanggal Export<br/>{export_d} &bull; {export_t} WIB", export_style))
    elements.append(Spacer(1, 20))

    summary_data = [
        ['Total Pemasukan',  f"Rp {int(data['total_pemasukan']):,}".replace(',', '.')],
        ['Total Pengeluaran', f"Rp {int(data['total_pengeluaran']):,}".replace(',', '.')],
        ['Laba Bersih',      f"Rp {int(data['laba_bersih']):,}".replace(',', '.')],
    ]
    t = Table(summary_data, colWidths=[200, 200])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.white),
        ('GRID',       (0, 0), (-1, -1), 1, colors.black),
        ('FONTNAME',   (0, 0), (-1, -1), 'Helvetica-Bold'),
        ('PADDING',    (0, 0), (-1, -1), 6),
    ]))
    elements.append(t)

    daily_rekap = _generate_daily_rekap(start_wib, end_wib, data)
    if daily_rekap:
        elements.append(Spacer(1, 20))
        elements.append(Paragraph("Rekap Harian", styles['Heading3']))
        rekap_table_data = [['Tanggal', 'Total Pemasukan', 'Total Pengeluaran', 'Laba Bersih']]
        for row in daily_rekap:
            rekap_table_data.append([
                row['tanggal'],
                f"Rp {int(row['pemasukan']):,}".replace(',', '.'),
                f"Rp {int(row['pengeluaran']):,}".replace(',', '.'),
                f"Rp {int(row['laba_bersih']):,}".replace(',', '.'),
            ])
        rt = Table(rekap_table_data, colWidths=[120, 130, 130, 130])
        rt.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
            ('GRID',       (0, 0), (-1, -1), 1, colors.black),
            ('PADDING',    (0, 0), (-1, -1), 6),
        ]))
        elements.append(rt)

    if data['top_products']:
        elements.append(Spacer(1, 20))
        elements.append(Paragraph("5 Produk Terlaris", styles['Heading3']))
        prod_data = [['Nama Produk', 'Terjual', 'Total Penjualan']]
        for p in data['top_products']:
            prod_data.append([p['name'], str(p['qty']),
                               f"Rp {int(p['subtotal']):,}".replace(',', '.')])
        pt = Table(prod_data, colWidths=[250, 80, 150])
        pt.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
            ('GRID',       (0, 0), (-1, -1), 1, colors.black),
            ('PADDING',    (0, 0), (-1, -1), 6),
        ]))
        elements.append(pt)

    on_page_func = _create_header_footer_watermark(start_wib, end_wib, active_biz)
    return elements, on_page_func


# ─────────────────────────────────────────────────────────────
#  ROUTE: Kirim ke Admin
# ─────────────────────────────────────────────────────────────
@laporan_bp.route('/kirim-ke-admin', methods=['POST'])
@login_required
def kirim_ke_admin():
    """
    Generate PDF laporan yang sama dengan export_pdf,
    simpan ke disk, dan catat pengiriman di tabel laporan_terkirim.
    Mengembalikan JSON untuk diproses AJAX di frontend.
    """
    from app.models.laporan_terkirim import LaporanTerkirim

    active_biz, err = _check_owner_and_active_biz()
    if err:
        return jsonify({'status': 'error', 'message': 'Usaha tidak ditemukan.'}), 400

    periode  = request.form.get('periode', 'bulan_ini')
    start_str = request.form.get('start_date')
    end_str   = request.form.get('end_date')

    start_wib, end_wib, periode = get_date_range(periode, start_str, end_str)
    formatted = format_periode_id(start_wib, end_wib, periode)
    data = get_report_data(active_biz, start_wib, end_wib, periode)

    # ── 1. Build PDF ke BytesIO ──────────────────────────────
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                            rightMargin=30, leftMargin=30,
                            topMargin=40, bottomMargin=50)
    elements, on_page_func = _build_pdf_elements(active_biz, start_wib, end_wib, periode, data)
    doc.build(elements, onFirstPage=on_page_func, onLaterPages=on_page_func)

    # ── 2. Simpan PDF ke INSTANCE_PATH/laporan_terkirim/ ─────
    save_dir = os.path.join(current_app.instance_path, 'laporan_terkirim')
    os.makedirs(save_dir, exist_ok=True)

    now_wib = datetime.now(_WIB)
    filename = (
        f"Laporan_{active_biz.id}_{now_wib.strftime('%Y%m%d_%H%M%S')}"
        f"_{start_wib.strftime('%Y%m%d')}_{end_wib.strftime('%Y%m%d')}.pdf"
    )
    file_path = os.path.join(save_dir, filename)

    with open(file_path, 'wb') as f:
        f.write(buffer.getvalue())

    # ── 3. Simpan record ke database ─────────────────────────
    record = LaporanTerkirim(
        business_id   = active_biz.id,
        sender_id     = current_user.id,
        periode       = periode,
        periode_label = formatted,
        start_date    = start_wib.date(),
        end_date      = end_wib.date(),
        submitted_at  = now_wib.astimezone(__import__('datetime').timezone.utc).replace(tzinfo=None),
        file_path     = os.path.join('laporan_terkirim', filename),
        status        = LaporanTerkirim.STATUS_BELUM,
    )
    db.session.add(record)
    db.session.commit()

    # ── 4. Buat WhatsApp Link untuk konfirmasi ───────────────
    from app.models.user import User
    from app.utils import normalize_whatsapp_number
    from urllib.parse import quote

    admin_user = User.query.filter_by(role='admin').first()
    admin_phone = normalize_whatsapp_number(admin_user.phone) if (admin_user and admin_user.phone) else ''

    message_text = (
        f"Halo Admin.\n\n"
        f"Saya telah mengirim laporan usaha melalui aplikasi Siwaruk.\n\n"
        f"Nama usaha:\n{active_biz.business_name}\n\n"
        f"Periode:\n{formatted}\n\n"
        f"Mohon untuk ditinjau.\n\n"
        f"Terima kasih."
    )
    encoded_text = quote(message_text)

    if admin_phone:
        wa_url = f"https://wa.me/{admin_phone}?text={encoded_text}"
    else:
        wa_url = f"https://wa.me/?text={encoded_text}"

    return jsonify({
        'status': 'success',
        'message': 'Laporan berhasil dikirim ke Admin.',
        'id': record.id,
        'periode_label': formatted,
        'business_name': active_biz.business_name,
        'wa_url': wa_url
    })

